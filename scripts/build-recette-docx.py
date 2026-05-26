#!/usr/bin/env python3
"""Génère MEATMASTER-FICHE-RECETTE-MOBILE.docx depuis le markdown source."""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = Path(__file__).resolve().parents[1]
MD_PATH = ROOT / "docs" / "MEATMASTER-FICHE-RECETTE-MOBILE.md"
OUT_DOCX = ROOT / "docs" / "MEATMASTER-FICHE-RECETTE-MOBILE.docx"
OUT_PDF = ROOT / "docs" / "MEATMASTER-FICHE-RECETTE-MOBILE.pdf"

INLINE_MD_RE = re.compile(r"\*\*(.+?)\*\*|`([^`]+)`")


def _apply_run_style(run, *, bold: bool = False, italic: bool = False, size: Pt | None = Pt(10), code: bool = False) -> None:
    if size is not None:
        run.font.size = size
    if bold:
        run.bold = True
    if italic:
        run.italic = True
    if code:
        run.font.name = "Consolas"


def append_inline_md(paragraph, text: str, *, base_bold: bool = False, base_italic: bool = False, size: Pt | None = Pt(11)) -> None:
    """Interprète **gras** et `code` dans un paragraphe Word."""
    if not text:
        return
    pos = 0
    for match in INLINE_MD_RE.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos : match.start()])
            _apply_run_style(run, bold=base_bold, italic=base_italic, size=size)
        if match.group(1) is not None:
            run = paragraph.add_run(match.group(1))
            _apply_run_style(run, bold=True, italic=base_italic, size=size)
        else:
            run = paragraph.add_run(match.group(2))
            _apply_run_style(run, bold=base_bold, italic=base_italic, size=size, code=True)
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        _apply_run_style(run, bold=base_bold, italic=base_italic, size=size)


def set_paragraph_rich(paragraph, text: str, *, base_bold: bool = False, base_italic: bool = False, size: Pt | None = Pt(11)) -> None:
    paragraph.text = ""
    append_inline_md(paragraph, text, base_bold=base_bold, base_italic=base_italic, size=size)


def add_rich_paragraph(doc: Document, text: str, style: str | None = None, **kwargs) -> None:
    p = doc.add_paragraph(style=style)
    set_paragraph_rich(p, text, **kwargs)


def set_cell_rich(cell, text: str, *, header: bool = False) -> None:
    p = cell.paragraphs[0]
    set_paragraph_rich(p, text, base_bold=header, size=Pt(10))
    if header:
        set_cell_shading(cell, "1F4E79")
        for run in p.runs:
            run.font.color.rgb = RGBColor(255, 255, 255)


def strip_inline_md(text: str) -> str:
    return INLINE_MD_RE.sub(lambda m: m.group(1) or m.group(2) or "", text)


def set_cell_shading(cell, fill: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    shading.set(qn("w:val"), "clear")
    cell._tc.get_or_add_tcPr().append(shading)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], header_fill: str = "1F4E79") -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        set_cell_rich(hdr[i], h, header=True)
    for ri, row in enumerate(rows):
        cells = table.rows[ri + 1].cells
        for ci, val in enumerate(row):
            set_cell_rich(cells[ci], val)
    doc.add_paragraph()


def add_meta_table(doc: Document, left: str, right: str) -> None:
    t = doc.add_table(rows=1, cols=2)
    t.style = "Table Grid"
    t.columns[0].width = Cm(4)
    t.columns[1].width = Cm(12)
    c0, c1 = t.rows[0].cells
    set_paragraph_rich(c0.paragraphs[0], left, base_bold=True, size=Pt(10))
    set_paragraph_rich(c1.paragraphs[0], right, size=Pt(10))
    doc.add_paragraph()


def add_ok_table(doc: Document) -> None:
    add_table(doc, ["OK", "KO", "Remarques"], [["☐", "☐", ""]])


def parse_md_to_docx(md_text: str, doc: Document) -> None:
    lines = md_text.splitlines()
    i = 0
    in_table = False
    table_rows: list[list[str]] = []

    def flush_table() -> None:
        nonlocal table_rows, in_table
        if not table_rows:
            in_table = False
            return
        headers = [c.strip() for c in table_rows[0]]
        body = [[c.strip() for c in row] for row in table_rows[2:]] if len(table_rows) > 2 else []
        if headers and headers[0].startswith("---"):
            in_table = False
            table_rows = []
            return
        add_table(doc, headers, body)
        table_rows = []
        in_table = False

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith("|") and "|" in stripped[1:]:
            if not in_table:
                in_table = True
                table_rows = []
            table_rows.append([c.strip() for c in stripped.strip("|").split("|")])
            i += 1
            continue
        if in_table:
            flush_table()

        if stripped == "---":
            doc.add_paragraph()
            i += 1
            continue

        if stripped.startswith("# "):
            p = doc.add_heading(strip_inline_md(stripped[2:].strip()), level=0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1
            continue

        if stripped.startswith("## "):
            doc.add_heading(strip_inline_md(stripped[3:].strip()), level=1)
            i += 1
            continue

        if stripped.startswith("### "):
            doc.add_heading(strip_inline_md(stripped[4:].strip()), level=2)
            i += 1
            continue

        label_val = re.match(r"^\*\*(.+?)\*\*:?\s*(.*)$", stripped)
        if label_val and not stripped.startswith("###"):
            add_meta_table(doc, label_val.group(1).strip(), label_val.group(2).strip())
            i += 1
            continue

        if stripped.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.5)
            set_paragraph_rich(p, stripped[2:].strip(), base_italic=True, size=Pt(10))
            i += 1
            continue

        if stripped == "**Actions à réaliser**" or stripped.startswith("**Actions à réaliser"):
            doc.add_paragraph("Actions à réaliser").runs[0].bold = True
            i += 1
            continue

        if stripped == "**Résultat attendu**":
            doc.add_paragraph("Résultat attendu").runs[0].bold = True
            i += 1
            continue

        if re.match(r"^\d+\.\s", stripped):
            add_rich_paragraph(doc, stripped, style="List Number")
            i += 1
            continue

        if stripped.startswith("- "):
            add_rich_paragraph(doc, stripped[2:].strip(), style="List Bullet")
            i += 1
            continue

        if stripped == "| OK | KO | Remarques |":
            add_ok_table(doc)
            i += 1
            continue

        if stripped.startswith("|") and "OK" in stripped:
            i += 1
            continue

        if stripped and not stripped.startswith("\\"):
            add_rich_paragraph(doc, stripped)
        i += 1

    if in_table:
        flush_table()


def main() -> int:
    out_docx = OUT_DOCX
    out_pdf = OUT_PDF
    if len(sys.argv) > 1:
        out_docx = Path(sys.argv[1])
        out_pdf = out_docx.with_suffix(".pdf")

    if not MD_PATH.exists():
        print(f"Missing {MD_PATH}", file=sys.stderr)
        return 1

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    md_text = MD_PATH.read_text(encoding="utf-8")
    # Skip yaml front matter
    if md_text.startswith("---"):
        end = md_text.find("\n---", 3)
        if end != -1:
            md_text = md_text[end + 4 :]

    parse_md_to_docx(md_text, doc)
    doc.save(out_docx)
    print(f"Created: {out_docx}")

    try:
        from docx2pdf import convert

        convert(str(out_docx), str(out_pdf))
        print(f"Created: {out_pdf}")
    except Exception as exc:
        print(f"PDF skipped (Word required): {exc}", file=sys.stderr)

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
